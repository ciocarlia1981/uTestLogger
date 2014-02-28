import xlrd
import re
from jinja2 import Environment, FileSystemLoader
import pandas
import datetime
import docx

book = xlrd.open_workbook("testDataSheet2.xls")

backgroundSheetName = "Background"

backgroundSheet = book.sheet_by_name(backgroundSheetName)

headerRow = 2

#print taskSheet.ncols



def getBackgroundSheet():
    pass


def getParticipants():
    participants = []
    for rowNumber in range(2, backgroundSheet.nrows):
        testDate = backgroundSheet.cell_value(rowNumber, 2)
        if testDate:
            newParticipant = {}
            for colNumber, colValue in enumerate(backgroundSheet.row_values(1)):
                #print colValue
                if colValue != "":
                    if backgroundSheet.cell_type(rowNumber, colNumber) == 1: # text
                        newParticipant[colValue] = backgroundSheet.cell_value(rowNumber, colNumber)
                    elif backgroundSheet.cell_type(rowNumber, colNumber) == 2: # float
                        newParticipant[colValue] = backgroundSheet.cell_value(rowNumber, colNumber)
                    elif backgroundSheet.cell_type(rowNumber, colNumber) == 3: # date
                        newParticipant[colValue] = backgroundSheet.cell_value(rowNumber, colNumber)
                    else:
                        newParticipant[colValue] = backgroundSheet.cell_value(rowNumber, colNumber)
                    newParticipant[colValue] = backgroundSheet.cell_value(rowNumber, colNumber)
                    #print "newParticipant[",colValue,"]: ", newParticipant[colValue]

            #print "New: ", newParticipant
            participants.append(newParticipant)
    participants = pandas.DataFrame(participants)
    return participants


def getAllJobTitles():
    jobTitles = []
    for participant in getParticipants():
        jobTitles.append(participant["Current Job Title"])

    jobTitles = set(jobTitles)
    return jobTitles


useErrorTypes = []

useErrors = []
uesByP = []


def getTasks():
    tasks = []
    for sheet in book.sheets():
        try:
            taskName = sheet.cell(0, 0).value

            if taskName[:4] in ("Task", "task"):
                description = taskName.split(".")[1][1:]
                number = taskName.split(".")[0][4:].replace(" ","")

                numberOfParticipantsWhoPerfTask = 0

                print "----"
                print description
                print number
            #numParticipantsWhoPerformedTask
            #Overall pass rate
            #Safety related pass rate
            #Safety related failure types - Use Error / Assists
            #Non-safety-related failure types - Use Error / Assists

            tasks.append(
                {"name:", taskName,
                }
            )
        except Exception as error:
            print error

    print tasks

getTasks()

def getUseErrors():
    global useErrors
    global uesByP
    for sheet in book.sheets():
        for rownum in range(1, sheet.nrows):
            pt = []
            row_values = sheet.row_values(rownum)

            for colNumber, value in enumerate(row_values):
                try:
                    searchResult = re.search('\[[uU][eE]\][^[]*', value)
                    if searchResult is not None:
                        if sheet.cell(0, 0).value[0] == "T" and sheet.cell(0, 1).value is not u'P#':
                            newUseError = {}
                            newUseError["task"] = sheet.cell(0, 0).value
                            newUseError["participant"] = sheet.cell(rownum, 0).value
                            newUseError["description"] = searchResult.group(0)
                            newUseError["useErrorName"] = sheet.cell(1, colNumber).value
                            useErrors.append(newUseError)

                except Exception as e:
                    pass
                    #print "ERROR: ", str(e)
                    #print ""

    useErrors = pandas.DataFrame(useErrors)
    useErrors = useErrors[useErrors.participant != u'P#']
    uesByP = useErrors.groupby('participant')
    print "useErrors:", useErrors


def getUseErrorTypes():
    useErrorTypes = []
    for useError in getUseErrors():
        useErrorTypes.append(useError["Use Error"])
    useErrorTypes = set(useErrorTypes)
    return useErrorTypes


#TODO : something news

#TODO: make the function more generic. Maybe indicate what data gets returned
def getParticipantSummary_byCategory(category, participants):
    summary = []
    subcategories = participants[category].unique()
    for subcategory in subcategories:
        categoryDetails = participants[ participants[category] == subcategory]
        hospitalUnits = categoryDetails['What department do you work in?'].unique()
        numMales = len(categoryDetails[categoryDetails['Gender'] == "M"])
        numFemales = len(categoryDetails[categoryDetails['Gender'] == "F"])
        avgAge = categoryDetails["Age"].mean()
        minAge = categoryDetails["Age"].min()
        maxAge = categoryDetails["Age"].max()

        summary.append({
            "title": subcategory,
            "numberP": len(categoryDetails),
            "hospitalUnits": hospitalUnits,
            "numMales": numMales,
            "numFemales": numFemales,
            "avgAge": avgAge,
            "minAge": minAge,
            "maxAge": maxAge
        })
    return summary

def getTaskPerformanceSummary():
    tasks = getTasks()


env = Environment(loader=FileSystemLoader("C:/Users/Jonathan/fun/uTestLogger"))
template = env.get_template('testTemplate.html')

participants = getParticipants()

jobTitles = participants['Current Job Title'].unique()

participantSummary = getParticipantSummary_byCategory('Current Job Title', participants)


templateVars = {
    "totalNumParticipants": len(participants),
    "jobTitles": participantSummary
}

outputText = template.render(templateVars)

print outputText

#getUseErrorTypes()

getUseErrors()

#getParticipants()

# regex for finding use errors: \[[uU][eE]\][^[]*
'''
for rownum in range(1, taskSheet.nrows):
    pt = []
    row_values = taskSheet.row_values(rownum)
    
    for value in row_values:
        try:
            searchResult = re.search('\[[uU][eE]\][^[]*', value)
            if searchResult is not None:
                print "---"
                print "Task: ", taskSheet.cell(0, 0).value
                print "Participant: ", taskSheet.cell(rownum, 0).value
                print searchResult.group(0)
        except Exception:
            doNothing = ""
            #print ""
        
        '''



#for colNum in range(1, taskSheet.ncols):
#   print taskSheet.row_values(headerRow)[colNum]
        
    
    
    


